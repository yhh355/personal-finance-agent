"""Build the polished Word submission from the version-controlled REPORT.md."""

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "REPORT.md"
OUTPUT = ROOT / "Personal_Finance_Agent_Final_Report.docx"


def set_run_font(run, name: str = "Aptos") -> None:
    """Apply a font consistently across Word renderers."""
    run.font.name = name
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)


def add_page_number(paragraph) -> None:
    """Add a dynamic PAGE field to a footer paragraph."""
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08

    for name, size in (("Title", 20), ("Heading 1", 14), ("Heading 2", 12)):
        style = doc.styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(31, 78, 121)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(14 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)

    subtitle = doc.styles.add_style("Report Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    subtitle.font.name = "Aptos"
    subtitle.font.size = Pt(11)
    subtitle.font.color.rgb = RGBColor(89, 89, 89)
    subtitle.paragraph_format.space_after = Pt(17)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(title.add_run("Personal Finance Statement Agent Final Report"), "Aptos Display")
    title.runs[0].font.color.rgb = RGBColor(31, 78, 121)

    meta = doc.add_paragraph(style="Report Subtitle")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(meta.add_run("PE6201 Emerging AI Technologies | End of Course Project"))

    for raw_line in SOURCE.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("# "):
            continue
        if line.startswith("## "):
            paragraph = doc.add_paragraph(style="Heading 1")
            set_run_font(paragraph.add_run(line[3:]))
            continue
        paragraph = doc.add_paragraph(style="Normal")
        set_run_font(paragraph.add_run(line))

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("PE6201 | Personal Finance Statement Agent | Page ")
    set_run_font(footer_run)
    footer_run.font.size = Pt(9)
    add_page_number(footer)

    doc.core_properties.title = "Personal Finance Statement Agent Final Report"
    doc.core_properties.subject = "PE6201 End of Course Project"
    doc.core_properties.author = ""
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
